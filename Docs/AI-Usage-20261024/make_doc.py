# -*- coding: utf-8 -*-
import pathlib
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.enum.text import WD_BREAK

OUT = pathlib.Path(r"D:\灯影山海\Docs\AI-Usage-20261024")
OUT.mkdir(parents=True, exist_ok=True)
DOC = OUT / "灯影山海_人工智能工具使用情况说明.docx"

BODY_CN, HEAD_CN = "宋体", "黑体"

def set_run_font(run, cn=BODY_CN, size=10.5, bold=False, color=None):
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.bold = bold
    if color: run.font.color.rgb = RGBColor(*color)
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts'); rPr.append(rFonts)
    rFonts.set(qn('w:eastAsia'), cn)
    rFonts.set(qn('w:ascii'), "Times New Roman")
    rFonts.set(qn('w:hAnsi'), "Times New Roman")

def para(text="", size=10.5, bold=False, cn=BODY_CN, align="left", space_before=0, space_after=4, indent_chars=0, line=1.5, color=None):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(space_before); pf.space_after = Pt(space_after)
    pf.line_spacing = line
    if indent_chars: pf.first_line_indent = Pt(size * indent_chars)
    pf.alignment = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER,
                    "justify": WD_ALIGN_PARAGRAPH.JUSTIFY, "right": WD_ALIGN_PARAGRAPH.RIGHT}[align]
    if text:
        r = p.add_run(text); set_run_font(r, cn=cn, size=size, bold=bold, color=color)
    return p

def heading(text, new_page=False):
    p = para(text, size=12, bold=True, cn=HEAD_CN, space_before=8, space_after=5)
    p.paragraph_format.keep_with_next = True
    if new_page:
        p.paragraph_format.page_break_before = True
    return p

def shade(cell, hexcolor="F2F2F2"):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd'); shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hexcolor); tcPr.append(shd)

def cell_text(cell, text, bold=False, size=9.5, align="left", cn=BODY_CN):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(1); p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.line_spacing = 1.15
    p.alignment = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER,
                   "justify": WD_ALIGN_PARAGRAPH.JUSTIFY}[align]
    for i, line in enumerate(str(text).split("\n")):
        if i: p = cell.add_paragraph(); p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(1); p.paragraph_format.line_spacing = 1.15
        r = p.add_run(line); set_run_font(r, cn=cn, size=size, bold=bold)

def table(rows, widths, header=True, header_fill="EDE9E1"):
    t = doc.add_table(rows=len(rows), cols=len(widths))
    t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for r_i, row in enumerate(rows):
        for c_i, val in enumerate(row):
            cell = t.cell(r_i, c_i); cell.width = Cm(widths[c_i])
            is_head = header and r_i == 0
            cell_text(cell, val, bold=is_head, align="center" if (is_head or c_i == 0 and len(widths) > 2) else "left")
            if is_head: shade(cell, header_fill)
    for row in t.rows:
        for c_i, cell in enumerate(row.cells):
            cell.width = Cm(widths[c_i])
        trPr = row._tr.get_or_add_trPr()
        cant = OxmlElement('w:cantSplit'); trPr.append(cant)
        if r_i == 0:
            trPr.append(OxmlElement('w:tblHeader'))
    # keep the whole table on one page when it fits: all but the last row keep with next
    for r_i, row in enumerate(t.rows[:-1]):
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.keep_with_next = True
    return t

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(2.6)
sec.top_margin = Cm(2.2); sec.bottom_margin = Cm(2.15)

normal = doc.styles["Normal"]; normal.font.name = "Times New Roman"; normal.font.size = Pt(10.5)
normal.element.rPr.rFonts.set(qn('w:eastAsia'), BODY_CN)

para("2026年第十四届全国大学生数字媒体科技作品及创意竞赛", size=12, bold=True, cn=HEAD_CN, align="center", space_after=2)
para("人工智能工具使用情况说明", size=16, bold=True, cn=HEAD_CN, align="center", space_after=6)
para("作品《灯影山海》", size=11, bold=True, align="center", space_after=2)
para("（按赛事匿名化要求编制，不填写学校、团队成员与指导教师信息）", size=9, align="center", space_after=12, color=(0x59,0x59,0x59))

heading("一、作品基本信息与AI使用状态")
table([
    ["项目", "填写内容"],
    ["作品名称", "灯影山海"],
    ["作品编号（如有）", "以报名系统分配编号为准"],
    ["参赛赛道／赛项", "虚幻引擎（UE）作品创作赛项"],
    ["作品形态", "Unreal Engine 5.8 实时影戏解谜 Demo，单主地图，Windows 平台可执行程序"],
    ["AI使用情况", "已使用（编程与建模辅助、文档与脚本起草；创意、玩法判定、关卡与美术定稿由学生完成）"],
], widths=[3.6, 12.2])

def bullet(text, size=10.5):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.left_indent = Cm(0.72); pf.first_line_indent = Cm(-0.42)
    pf.space_before = Pt(0); pf.space_after = Pt(2); pf.line_spacing = 1.4
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p.add_run("·  " + text); set_run_font(r, size=size)
    return p

heading("二、主要AI工具及使用情况")
para("仅列对最终成果具有实际影响的主要工具；日常资料查询不在此列。", size=9, color=(0x59,0x59,0x59), space_after=4)
table([
    ["工具／模型", "使用环节", "AI主要用途或产出", "学生后续处理"],
    ["AI 编程与建模助手\n（Codex 客户端环境，GPT 系列代码模型）",
     "Blender 程序化建模、材质与贴图生成、工具脚本",
     "三具影偶的分件提取、真实镂空与 0.65 cm 皮厚网格、骨骼与姿态导出脚本；材质贴图与预览渲染脚本；FBX 往返自检脚本",
     "逐图人工检查（标注残留、轮廓毛边、孔洞），删除失败版本并保留回退；模型比例、关节与材质最终由人工审定"],
    ["AI 编程与审查助手\n（DeepSeek Harness 客户端环境，DeepSeek 系列模型）",
     "C++ 逻辑修复、整体审查、打包与验证脚本",
     "换偶与昼夜眼睑逻辑补丁草案；HUD 文案与音效去重分析；打包、冷重载、四幕截图与放映胶片段的自动化脚本；文档起草",
     "逐条在正式关卡与打包版本中实测，核对探针数据后决定采纳或回退；判定参数与体验节奏不改由 AI 决定"],
], widths=[3.4, 3.4, 5.4, 3.6])
para("说明：如团队另使用过图像或语音生成工具，可据实补充一行；本作品未使用语音合成，评审视频为静音版。", size=9, color=(0x59,0x59,0x59), space_before=3)

heading("三、代表性AI使用过程", new_page=True)
sub1 = para("（一）三具影偶：从平面概念图到可投影的实体皮影", size=10.5, bold=True, space_before=4, space_after=3)
sub1.paragraph_format.keep_with_next = True
table([
    ["说明项目", "具体内容"],
    ["使用环节及目的", "为第二至第四幕补充精卫、刑天、烛龙三具影偶，使每幕有独立角色，而非四幕共用同一具皮影。"],
    ["主要工具／模型", "Codex 环境下的 AI 编程与建模助手；Blender 4.5 无头模式；Python（OpenCV、Shapely、triangle）。"],
    ["主要输入或素材来源", "团队自绘的三张正交概念图（Jingwei／Xingtian／Zhulong），含分件、关节与杆位标注；无第三方模型或贴图。"],
    ["AI主要输出", "按概念图像素坐标提取 17／20／6 个分件轮廓，保留 99／33／32 个真实透光孔，挤出 0.65 cm 皮厚并生成 UV、刚性骨骼（17／14／8 根）与三套姿态；导出 FBX 与 2K BaseColor／NormalDX／ORM 贴图。"],
    ["学生筛选、修改或再创作", "逐张实机图核对：删除混入的尺寸线与箭头、修正翼尖与龙身边缘、补齐烛龙可开合眼睑、按实际翼展修正比例记录；Run1 整批保留为回退，最终采用 Run2。"],
    ["最终成果中的应用位置", "第二幕精卫、第三幕刑天、第四幕烛龙；资源位于作品 Puppets／Dev 目录，专用检查地图 L_ThreePuppetsReview 保留 12 张 Movable 光源投影核验图。"],
    ["可核验材料（如有）", "模型验证记录（三角面、骨骼、包围盒、FBX 冷读）、UE 导入记录与 12 张投影截图、冷重载回执。"],
], widths=[3.4, 12.4])

sub2 = para("（二）四幕换偶与烛龙昼夜：AI 草案经人工实测纠正", size=10.5, bold=True, space_before=8, space_after=3)
sub2.paragraph_format.page_break_before = True
sub2.paragraph_format.keep_with_next = True
table([
    ["说明项目", "具体内容"],
    ["使用环节及目的", "让四幕在正式关卡中真正切换影偶，并使烛龙的眼睑随昼夜变化开合。"],
    ["主要工具／模型", "AI 编程与审查助手（DeepSeek Harness 环境）；UE 5.8 C++ 与编辑器探针脚本。"],
    ["主要输入或素材来源", "现有 Puzzle Director、Puppet Actor 源码；正式主地图与打包包；人工记录的探针结果。"],
    ["AI主要输出", "CastMesh 单节点动画、按章切换显示姿态、眼睑缩放分支等补丁草案；诊断脚本与结论摘要。"],
    ["学生筛选、修改或再创作", "实测发现草案首版在正式关卡仍停在 Blueprint 动画模式，人工补充探针定位后改为显式切单节点并重设位置；确认夜伏成功时不应提前睁眼，重排昼夜判定；随后重新打包并复测四幕。"],
    ["最终成果中的应用位置", "四幕换偶与第四幕昼夜表现；正式主地图 L_DY_Main 与 Shipping 包。"],
    ["可核验材料（如有）", "探针 JSON、打包 identity（源图与可执行文件 SHA256）、四幕实机截图与录屏。"],
], widths=[3.4, 12.4])

heading("四、最终成果中的AI生成或辅助内容")
table([
    ["内容／成果", "AI生成或辅助部分", "学生修改与整合", "最终使用位置"],
    ["三具影偶的网格、UV、骨骼与贴图", "轮廓提取、镂空、挤出、绑骨与贴图烘焙脚本由 AI 生成", "人工去掉图稿标注、修轮廓、定厚度与比例、逐张验收", "精卫／刑天／烛龙角色资产与三套姿态"],
    ["部分 C++ 逻辑与工具脚本", "补丁与脚本草案", "人工复核、编译、打包回归，失败项回退", "换偶、眼睑、HUD 文案、音效去重、打包与验证脚本"],
    ["文字与图示材料", "文档初稿、表格与说明文字", "人工核对数据与口径，改写为提交版本", "本说明、交付记录与验证报告（非作品本体）"],
    ["评审视频的剪辑脚本", "分镜与转场脚本、字幕排版脚本", "人工挑帧、定稿与复核时长", "参赛评审视频（非实时作品本体）"],
], widths=[4.0, 4.4, 4.4, 3.4])

heading("五、学生原创贡献说明")
para("作品的核心创意与实施方案由学生提出并定稿，AI 承担的是可替换的实现辅助。", indent_chars=2)
for line in [
    "主题与叙事：确定以《山海经》夸父逐日与弃杖化林为母题的原创改编方向，形成“久未开演—未完逐日—玩家复演”的最小故事，设定四幕（夸父逐日／精卫填海／刑天舞干戚／烛龙昼夜）、题签与字幕文本。",
    "玩法与判定：确立“光源即玩法、影子即角色”，以灯位、偶位与保持时间三参数的区间判定作为唯一主判定，视觉阴影与逻辑判定分离；三次互动沿“发现规律—协同控制—综合运用”递进，第三次只组合已学能力。",
    "关卡与美术方向：规划单主地图的入场灯径、幕后操作区、左右台翼、简化台前与不可探索远景；确定影窗、台口、壁柱、八字纸灯、锣鼓曲面网格与观众席的构成，并对 AI 产出一项项取舍（如三次树冠重建失败后回滚、门脸格窗方案删除）。",
    "工程与验收：完成 UE 项目配置、蓝图接线、灯光与相机布置、打包与发布；对每轮改动做真机冒烟、四幕通关、截图与录屏核验，保留回退基线。",
]:
    bullet(line)
para("AI 与学生工作边界：AI 不参与创意取舍与验收判定；凡 AI 产出的代码、网格、贴图与文字，均须经学生逐项检查并在正式关卡或打包版本中实测后才进入作品。未通过实测的草案（如首版换偶补丁、程序化树冠、门脸格窗）均已回退，不作为成果保留。", indent_chars=2, space_before=4)

heading("六、代表性过程材料（按需）")
para("按赛事要求，可在参赛打包文件中附 1—3 组必要材料：三具影偶的分件与镂空对比图、换偶探针前后对照、四幕实机截图。材料已脱敏，不含账号与个人信息；不提交完整对话记录。", indent_chars=2)

heading("七、资源来源与合规说明")
for line in [
    "引擎与基础资源：Unreal Engine 5.8.3，按 Epic Games 许可用于本作品；引擎自带基础形状与材质仅作占位与调试。",
    "字体：Noto Serif SC、Noto Sans SC，SIL Open Font License 1.1，用于字幕、界面与文档排版。",
    "图像与模型：主视觉、概念图与场景资产为团队自绘或团队自制程序化生成；三具影偶基于团队概念图重建，未使用第三方模型库。",
    "音频：锣鼓与环境音按项目《第三方资源说明》逐项登记，均为免费／开源许可或团队自制，无付费素材。",
    "AI 生成内容：均在人工筛选与修改后使用，未将 AI 输出直接表述为学生独立原创；使用的工具与环节见本说明第二、三节。",
]:
    bullet(line)

heading("八、填写确认")
para("本团队确认：本说明真实反映人工智能工具在《灯影山海》中的使用情况、学生自主完成的工作与最终成果状态；未虚构使用过程，未将 AI 或第三方成果表述为学生独立原创。相关真实性与知识产权承诺按本届《参赛承诺书》执行。", indent_chars=2)
para("填表日期：2026 年 10 月 24 日", align="right", space_before=8)

footer_p = sec.footer.paragraphs[0]
footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
def add_field(par, instr):
    r = par.add_run(); set_run_font(r, size=9)
    fld = OxmlElement('w:fldSimple'); fld.set(qn('w:instr'), instr)
    inner_r = OxmlElement('w:r'); inner_t = OxmlElement('w:t'); inner_t.text = "1"
    inner_r.append(inner_t); fld.append(inner_r)
    par._p.append(fld)
r0 = footer_p.add_run("第 "); set_run_font(r0, size=9)
add_field(footer_p, "PAGE")
r1 = footer_p.add_run(" 页 / 共 "); set_run_font(r1, size=9)
add_field(footer_p, "NUMPAGES")
r2 = footer_p.add_run(" 页"); set_run_font(r2, size=9)

doc.save(str(DOC))
print("SAVED", DOC, DOC.stat().st_size)
