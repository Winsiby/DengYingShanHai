# -*- coding: utf-8 -*-
import os, urllib.parse, urllib.request, pathlib
outdir = pathlib.Path(r"D:\灯影山海\Docs\AI-Usage-20261024")
outdir.mkdir(parents=True, exist_ok=True)
name = "2026年第十四届全国大参赛团队数字媒体科技作品及创意竞赛人工智能工具使用情况说明.docx"
url = "https://www.cmit.cn/uploadfile/2026/1015/" + urllib.parse.quote(name)
dst = outdir / "官方模板.docx"
req = urllib.request.Request(url, headers={"User-Agent":"Mozilla/5.0"})
with urllib.request.urlopen(req, timeout=60) as r:
    data = r.read()
dst.write_bytes(data)
print("bytes", len(data))

from docx import Document
d = Document(str(dst))
print("PARAGRAPHS", len(d.paragraphs), "TABLES", len(d.tables))
for i,p in enumerate(d.paragraphs):
    if p.text.strip(): print(f"P{i}|{p.style.name}|{p.text.strip()[:120]}")
for ti,t in enumerate(d.tables):
    print(f"--- TABLE {ti} rows={len(t.rows)} cols={len(t.columns)}")
    for ri,row in enumerate(t.rows[:8]):
        print("  ", ri, " || ".join(c.text.strip().replace("\n"," / ")[:60] for c in row.cells))
